"""
File Reader Agent
=================
Agent 1 of 9 in the Agentic AI Pipeline.

Detects file types (CSV, Excel, PDF, DOCX) and reads their contents into a raw structure.
"""

import os
from pypdf import PdfReader
from docx import Document
import pandas as pd
import io

class FileReaderAgent:
    """Agent responsible for reading raw files of different formats."""

    def read_file(self, file_bytes: bytes, filename: str) -> dict:
        """
        Detect format and read the file content.
        
        Returns:
            dict: {
                "format": str ("csv", "xlsx", "pdf", "docx"),
                "raw_content": str or DataFrame,
                "success": bool,
                "error": str
            }
        """
        ext = os.path.splitext(filename)[1].lower()
        try:
            if ext == '.csv':
                # Try semicolon first, then comma
                file_stream = io.BytesIO(file_bytes)
                try:
                    df = pd.read_csv(file_stream, sep=';')
                    if df.shape[1] <= 1:
                        file_stream.seek(0)
                        df = pd.read_csv(file_stream, sep=',')
                except Exception:
                    file_stream.seek(0)
                    df = pd.read_csv(file_stream, sep=',')
                return {"format": "csv", "raw_content": df, "success": True, "error": None}

            elif ext in ['.xlsx', '.xls']:
                df = pd.read_excel(io.BytesIO(file_bytes))
                return {"format": "xlsx", "raw_content": df, "success": True, "error": None}

            elif ext == '.pdf':
                reader = PdfReader(io.BytesIO(file_bytes))
                text_content = ""
                for page in reader.pages:
                    text_content += page.extract_text() or ""
                return {"format": "pdf", "raw_content": text_content, "success": True, "error": None}

            elif ext == '.docx':
                doc = Document(io.BytesIO(file_bytes))
                text_content = []
                # Extract paragraph text
                for p in doc.paragraphs:
                    if p.text:
                        text_content.append(p.text)
                # Extract tables text if any
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text for cell in row.cells]
                        text_content.append(" | ".join(row_text))
                return {"format": "docx", "raw_content": "\n".join(text_content), "success": True, "error": None}

            else:
                return {
                    "format": "unknown",
                    "raw_content": None,
                    "success": False,
                    "error": f"Unsupported file type: {ext}"
                }

        except Exception as e:
            return {
                "format": "error",
                "raw_content": None,
                "success": False,
                "error": f"Parsing failed: {str(e)}"
            }
